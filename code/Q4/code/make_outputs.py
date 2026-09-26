# -*- coding: utf-8 -*-
"""Aggregate ALNS-Q4 runs into AGENTS.md interface outputs.

Reads results/runs/k{2,3}_seed*/solution.json + pareto_archive.json, picks the
lexicographic best per K across seeds, merges per-K archives (dedup +
non-dominated re-filter), and writes:
  partition_result.csv   (AGENTS 4.5)
  resource_gap.csv       (AGENTS 4.5, per-group rows + TOTAL rows with reasons)
  pareto_archive.json    (merged, with frozen-plan invariants per AGENTS 3.5)
  global_audit.json      (8.3 audits + enumeration gap-0 verification)
  Q4_结果提交.xlsx        (optional, fills template sheet Q4_分区配置)
"""
import csv
import glob
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from q4_core import (INVENTORY, RESOURCE_TYPES, Context, evaluate_partition,   # noqa: E402
                     partition_label)
from q4_audit import audit_partition                                            # noqa: E402
from alns_q4 import archive_key, _dominates                                     # noqa: E402


def load_runs(runs_dir):
    runs = {2: [], 3: []}
    for path in sorted(glob.glob(os.path.join(runs_dir, 'k*_seed*',
                                              'solution.json'))):
        with open(path, encoding='utf-8') as f:
            sol = json.load(f)
        arch_path = os.path.join(os.path.dirname(path), 'pareto_archive.json')
        arch = None
        if os.path.exists(arch_path):
            with open(arch_path, encoding='utf-8') as f:
                arch = json.load(f)
        runs[sol['k']].append((sol, arch, os.path.dirname(path)))
    return runs


def pick_best(ctx, runs):
    chosen = {}
    for k in (2, 3):
        evs = []
        for sol, _, _ in runs[k]:
            ev = evaluate_partition(ctx, sol['best']['assignment'], k)
            evs.append((sol['seed'], ev))
        chosen[k] = min(evs, key=lambda x: archive_key(x[1]))
    return chosen


def merge_archive(ctx, runs, k):
    seen, entries = set(), []
    for sol, arch, _ in runs[k]:
        if not arch:
            continue
        for e in arch['entries']:
            key = tuple(e['assignment'])
            if key in seen:
                continue
            seen.add(key)
            entries.append(evaluate_partition(ctx, e['assignment'], k))
    front = [ev for ev in entries
             if not any(o is not ev and _dominates(o, ev) for o in entries)]
    front.sort(key=archive_key)
    return front


def gap_reason(ev, rt):
    if ev['shortage'][rt] == 0:
        return ''
    parts = []
    for g in ev['groups']:
        n = g['counts'][rt]
        if n:
            parts.append('G%d需%d(峰值t≈%.0fs)' % (g['group'] + 1, n,
                                                  g['witnesses'][rt]))
    return ('；'.join(parts) + '，合计%d>库存%d，任务执行期间不得跨组调配'
            % (ev['totals'][rt], INVENTORY[rt]))


def write_interface(ctx, chosen, out_dir):
    with open(os.path.join(out_dir, 'partition_result.csv'), 'w',
              newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['partition_count', 'group_id', 'service_node',
                    'component_id', 'workload_time_s', 'workload_energy_kwh',
                    'box_mass_kg'])
        for k in (2, 3):
            ev = chosen[k][1]
            for g in ev['groups']:
                node_wt, node_we, node_wm = {}, {}, {}
                for t in ctx.trips:
                    if t['trip'] not in g['trips']:
                        continue
                    dur = t['return_time'] - t['start']
                    stop_mass = {s: sum(ctx.box_attrs[b]['mass'] for b in t['boxes']
                                        if ctx.box_attrs[b]['site'] == s)
                                 for s in t['route']}
                    tot = sum(stop_mass.values())
                    for s in t['route']:
                        frac = stop_mass[s] / tot if tot else 1.0 / len(t['route'])
                        node_wt[s] = node_wt.get(s, 0) + dur * frac
                        node_we[s] = node_we.get(s, 0) + t['energy'] * frac
                        node_wm[s] = node_wm.get(s, 0) + stop_mass[s]
                for s in sorted(node_wt):
                    w.writerow([k, 'G%d' % (g['group'] + 1), s,
                                ctx.site_to_component[s],
                                '%.1f' % node_wt[s], '%.4f' % node_we[s],
                                '%.1f' % node_wm[s]])

    with open(os.path.join(out_dir, 'resource_gap.csv'), 'w',
              newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['partition_count', 'group_id', 'resource_type',
                    'required_count', 'inventory_count', 'redundancy_count',
                    'shortage_count', 'shortage_reason'])
        for k in (2, 3):
            ev = chosen[k][1]
            for g in ev['groups']:
                for rt in RESOURCE_TYPES:
                    w.writerow([k, 'G%d' % (g['group'] + 1), rt,
                                g['counts'][rt], INVENTORY[rt], '', '', ''])
            for rt in RESOURCE_TYPES:
                w.writerow([k, 'TOTAL', rt, ev['totals'][rt], INVENTORY[rt],
                            ev['redundancy'][rt], ev['shortage'][rt],
                            gap_reason(ev, rt)])


def fill_template(ctx, chosen, out_dir, template_path):
    if not os.path.exists(template_path):
        return 'template not found, skipped: %s' % template_path
    try:
        import openpyxl
    except Exception:
        return 'openpyxl unavailable, skipped'
    wb = openpyxl.load_workbook(template_path)
    if 'Q4_分区配置' not in wb.sheetnames:
        wb.close()
        return 'sheet Q4_分区配置 not found, skipped'
    ws = wb['Q4_分区配置']
    row = 2
    for k in (2, 3):
        ev = chosen[k][1]
        for g in ev['groups']:
            sites = []
            for cid in g['components']:
                sites.extend(ctx.components[cid])
            ws.cell(row=row, column=1, value=k)
            ws.cell(row=row, column=2, value='G%d' % (g['group'] + 1))
            ws.cell(row=row, column=3, value=','.join(sorted(sites)))
            for j, rt in enumerate(RESOURCE_TYPES):
                ws.cell(row=row, column=4 + j, value=g['counts'][rt])
            row += 1
    out = os.path.join(out_dir, 'Q4_结果提交.xlsx')
    wb.save(out)
    wb.close()
    return out


def run_outputs(ctx, results_dir, template_path=None):
    runs_dir = os.path.join(results_dir, 'runs')
    runs = load_runs(runs_dir)
    chosen = pick_best(ctx, runs)

    write_interface(ctx, chosen, results_dir)

    archives = {}
    for k in (2, 3):
        front = merge_archive(ctx, runs, k)
        archives['k%d' % k] = [dict(assignment=list(ev['assignment']),
                                    label=partition_label(ctx, ev),
                                    score=list(ev['score']),
                                    shortage_total=ev['shortage_total'],
                                    balance_time=ev['balance_time'],
                                    redundancy_total=ev['redundancy_total'],
                                    relay_extra=ev['relay_extra'],
                                    within_inventory=ev['within_inventory'])
                               for ev in front]
    with open(os.path.join(results_dir, 'pareto_archive.json'), 'w',
              encoding='utf-8') as f:
        json.dump(dict(objective_vector=['shortage_total', 'balance_time',
                                         'redundancy_total', 'relay_extra'],
                       invariants=ctx.invariants(),
                       archives=archives),
                  f, ensure_ascii=False, indent=2)

    # enumeration cross-check (gap = 0 verification)
    enum_summary = {}
    enum_path = os.path.join(results_dir, 'enumeration_summary.json')
    if os.path.exists(enum_path):
        with open(enum_path, encoding='utf-8') as f:
            enum_summary = json.load(f)

    audit = {'invariants': ctx.invariants(),
             'global_counts': ctx.global_counts, 'inventory': dict(INVENTORY)}
    for k in (2, 3):
        seed, ev = chosen[k]
        a = audit_partition(ctx, ev)
        a['seed'] = seed
        a['partition'] = partition_label(ctx, ev)
        a['per_seed_best'] = [
            dict(seed=sol['seed'], label=partition_label(
                     ctx, evaluate_partition(ctx, sol['best']['assignment'], k)),
                 score=list(evaluate_partition(ctx, sol['best']['assignment'], k)['score']))
            for sol, _, _ in runs[k]]
        eb = enum_summary.get('k%d' % k, {}).get('lexicographic_best')
        if eb:
            a['enumeration_best'] = eb['partition']
            a['enumeration_score'] = eb['score']
            a['alns_equals_enumeration'] = (
                list(ev['score']) == [int(x) if i != 3 else x
                                      for i, x in enumerate(eb['score'])])
            a['gap_to_enumeration'] = 0 if a['alns_equals_enumeration'] else None
        audit['k%d' % k] = a

    template_note = None
    if template_path:
        template_note = fill_template(ctx, chosen, results_dir, template_path)
    audit['template_fill'] = template_note

    with open(os.path.join(results_dir, 'global_audit.json'), 'w',
              encoding='utf-8') as f:
        json.dump(audit, f, ensure_ascii=False, indent=2)

    manifest = dict(runs={str(k): [os.path.relpath(p, results_dir)
                                   for _, _, p in runs[k]] for k in (2, 3)},
                    chosen={str(k): dict(seed=chosen[k][0],
                                         label=partition_label(ctx, chosen[k][1]),
                                         score=list(chosen[k][1]['score']))
                            for k in (2, 3)})
    with open(os.path.join(results_dir, 'run_manifest.json'), 'w',
              encoding='utf-8') as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
    return audit


if __name__ == '__main__':
    root = os.path.abspath(os.path.join(os.path.dirname(__file__),
                                        '..', '..', '..'))
    results = os.path.abspath(os.path.join(os.path.dirname(__file__),
                                           '..', 'results'))
    ctx = Context(os.path.join(results, 'frozen_alns_q3_v2_seed_20260924'))
    tpl = os.path.join(root, '结果提交模板.xlsx')
    audit = run_outputs(ctx, results, os.path.abspath(tpl))
    for k in (2, 3):
        a = audit['k%d' % k]
        print('K=%d feasible=%s gap0=%s partition=%s'
              % (k, a['feasible'], a.get('alns_equals_enumeration'),
                 a['partition']))
