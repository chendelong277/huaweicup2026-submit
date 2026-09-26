from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import platform
import sys
import time
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.io import loadmat


G = 9.80665
ENERGY_MODEL_VERSION = "q1_energy_v1_equivalent_range_plus_climb"
GEOMETRY_MODEL_VERSION = "q1_dem_supercover_v1"


@dataclass(frozen=True)
class Node:
    node_id: str
    node_type: str
    lon: float
    lat: float
    ground_elevation_m: float


@dataclass(frozen=True)
class Vehicle:
    vehicle_type: str
    vehicle_name: str
    empty_mass_kg: float
    max_payload_kg: float
    max_volume_m3: float
    cruise_speed_mps: float
    empty_range_m: float
    full_range_m: float
    usable_energy_kwh: float
    reserve_fraction: float
    fixed_prep_s: float
    per_box_load_s: float
    base_handoff_s: float
    per_box_handoff_s: float
    climb_speed_mps: float
    descent_speed_mps: float
    climb_efficiency: float


@dataclass(frozen=True)
class Box:
    box_id: str
    service_node: str
    material_type: str
    mass_kg: float
    volume_m3: float
    is_first_batch: bool
    deadline_s: float | None
    expected_time_s: float | None
    priority: float


@dataclass(frozen=True)
class Leg:
    from_node: str
    to_node: str
    horizontal_distance_m: float
    max_dem_elevation_m: float
    cruise_altitude_m: float
    climb_m: float
    descent_m: float
    intersected_cells: int


def find_one(root: Path, filename: str) -> Path:
    matches = list(root.rglob(filename))
    if len(matches) != 1:
        raise FileNotFoundError(f"Expected one {filename!r} below {root}, found {len(matches)}")
    return matches[0]


def sheet_values(path: Path, sheet: str | int = 0) -> list[tuple[Any, ...]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    return list(ws.values)


def load_inputs(data_root: Path) -> tuple[dict[str, Node], dict[str, Vehicle], dict[str, list[Box]], dict[str, Any], list[Path]]:
    node_file = find_one(data_root, "调度中心与服务区.xlsx")
    vehicle_file = find_one(data_root, "运输无人机数据.xlsx")
    demand_file = find_one(data_root, "物资需求与配送时限.xlsx")
    dem_file = find_one(data_root, "镇龙乡及周边30米DEM.mat")

    nodes: dict[str, Node] = {}
    for row in sheet_values(node_file):
        if not row or not isinstance(row[0], str):
            continue
        node_id = row[0].strip()
        if node_id == "O01" or (node_id.startswith("S") and node_id[1:].isdigit()):
            nodes[node_id] = Node(
                node_id=node_id,
                node_type="depot" if node_id == "O01" else "service",
                lon=float(row[2]),
                lat=float(row[3]),
                ground_elevation_m=float(row[4]),
            )

    vehicles: dict[str, Vehicle] = {}
    for row in sheet_values(vehicle_file):
        if not row or row[0] not in {"A", "B", "C"} or not isinstance(row[3], (int, float)):
            continue
        vehicles[row[0]] = Vehicle(
            vehicle_type=str(row[0]),
            vehicle_name=str(row[1]),
            empty_mass_kg=float(row[2]),
            max_payload_kg=float(row[3]),
            max_volume_m3=float(row[4]),
            cruise_speed_mps=float(row[5]),
            empty_range_m=float(row[6]),
            full_range_m=float(row[7]),
            usable_energy_kwh=float(row[8]),
            reserve_fraction=float(row[9]) / 100.0,
            fixed_prep_s=float(row[10]),
            per_box_load_s=float(row[11]),
            base_handoff_s=float(row[12]),
            per_box_handoff_s=float(row[13]),
            climb_speed_mps=float(row[14]),
            descent_speed_mps=float(row[15]),
            climb_efficiency=float(row[16]),
        )

    boxes_by_service: dict[str, list[Box]] = defaultdict(list)
    rows = sheet_values(demand_file, "逐箱货箱清单")
    for row in rows[1:]:
        if not row or not isinstance(row[0], str):
            continue
        box = Box(
            box_id=str(row[0]),
            service_node=str(row[1]),
            material_type=str(row[2]),
            mass_kg=float(row[3]),
            volume_m3=float(row[4]),
            is_first_batch=str(row[5]).strip() == "是",
            deadline_s=None if row[6] is None else float(row[6]),
            expected_time_s=None if row[7] is None else float(row[7]),
            priority=float(row[8]),
        )
        boxes_by_service[box.service_node].append(box)

    dem = loadmat(dem_file)
    required = {"dem", "longitude", "latitude", "nodata"}
    if not required <= set(dem):
        raise ValueError(f"DEM mat is missing fields: {sorted(required - set(dem))}")
    if set(boxes_by_service) - set(nodes):
        raise ValueError("Demand contains service nodes absent from the node workbook")
    if set(vehicles) != {"A", "B", "C"}:
        raise ValueError("Vehicle table did not yield exactly A, B and C")
    return nodes, vehicles, dict(boxes_by_service), dem, [node_file, vehicle_file, demand_file, dem_file]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def local_xy_m(node: Node, origin: Node) -> tuple[float, float]:
    mean_lat = math.radians((node.lat + origin.lat) / 2.0)
    x = (node.lon - origin.lon) * 111320.0 * math.cos(mean_lat)
    y = (node.lat - origin.lat) * 111195.0
    return x, y


def horizontal_distance_m(a: Node, b: Node) -> float:
    mean_lat = math.radians((a.lat + b.lat) / 2.0)
    dx = (b.lon - a.lon) * 111320.0 * math.cos(mean_lat)
    dy = (b.lat - a.lat) * 111195.0
    return math.hypot(dx, dy)


def supercover_cells(x0: float, y0: float, x1: float, y1: float, ncols: int, nrows: int) -> list[tuple[int, int]]:
    """Enumerate every raster cell intersected by a line in boundary-based grid coordinates."""
    eps = 1e-12
    x0 = min(max(x0, eps), ncols - eps)
    x1 = min(max(x1, eps), ncols - eps)
    y0 = min(max(y0, eps), nrows - eps)
    y1 = min(max(y1, eps), nrows - eps)
    ix, iy = math.floor(x0), math.floor(y0)
    ex, ey = math.floor(x1), math.floor(y1)
    dx, dy = x1 - x0, y1 - y0
    step_x = 0 if abs(dx) < eps else (1 if dx > 0 else -1)
    step_y = 0 if abs(dy) < eps else (1 if dy > 0 else -1)
    inf = float("inf")
    t_delta_x = inf if step_x == 0 else abs(1.0 / dx)
    t_delta_y = inf if step_y == 0 else abs(1.0 / dy)
    next_x = (ix + 1.0) if step_x > 0 else float(ix)
    next_y = (iy + 1.0) if step_y > 0 else float(iy)
    t_max_x = inf if step_x == 0 else (next_x - x0) / dx
    t_max_y = inf if step_y == 0 else (next_y - y0) / dy
    cells: set[tuple[int, int]] = {(iy, ix)}
    guard = 0
    while (ix, iy) != (ex, ey):
        guard += 1
        if guard > ncols + nrows + 10:
            raise RuntimeError("Raster traversal guard triggered")
        if abs(t_max_x - t_max_y) <= 1e-12:
            nx, ny = ix + step_x, iy + step_y
            if 0 <= ix + step_x < ncols:
                cells.add((iy, ix + step_x))
            if 0 <= iy + step_y < nrows:
                cells.add((iy + step_y, ix))
            ix, iy = nx, ny
            t_max_x += t_delta_x
            t_max_y += t_delta_y
        elif t_max_x < t_max_y:
            ix += step_x
            t_max_x += t_delta_x
        else:
            iy += step_y
            t_max_y += t_delta_y
        if 0 <= ix < ncols and 0 <= iy < nrows:
            cells.add((iy, ix))
    return sorted(cells)


def dem_cells_for_leg(a: Node, b: Node, dem: dict[str, Any]) -> list[tuple[int, int]]:
    lon = np.asarray(dem["longitude"])[0]
    lat = np.asarray(dem["latitude"])[:, 0]
    dx = float(lon[1] - lon[0])
    dy = float(lat[0] - lat[1])
    x0 = (a.lon - float(lon[0])) / dx + 0.5
    x1 = (b.lon - float(lon[0])) / dx + 0.5
    y0 = (float(lat[0]) - a.lat) / dy + 0.5
    y1 = (float(lat[0]) - b.lat) / dy + 0.5
    return supercover_cells(x0, y0, x1, y1, len(lon), len(lat))


def build_leg(a: Node, b: Node, a_operation_altitude_m: float, b_operation_altitude_m: float, dem: dict[str, Any]) -> Leg:
    cells = dem_cells_for_leg(a, b, dem)
    raster = np.asarray(dem["dem"])
    nodata = float(np.asarray(dem["nodata"]).ravel()[0])
    values = np.array([float(raster[row, col]) for row, col in cells])
    if np.any(values == nodata) or np.any(~np.isfinite(values)):
        raise ValueError(f"DEM NoData encountered on {a.node_id}->{b.node_id}")
    max_dem = float(values.max())
    cruise = max_dem + 50.0
    return Leg(
        from_node=a.node_id,
        to_node=b.node_id,
        horizontal_distance_m=horizontal_distance_m(a, b),
        max_dem_elevation_m=max_dem,
        cruise_altitude_m=cruise,
        climb_m=max(0.0, cruise - a_operation_altitude_m),
        descent_m=max(0.0, cruise - b_operation_altitude_m),
        intersected_cells=len(cells),
    )


def equivalent_range_m(vehicle: Vehicle, payload_kg: float) -> float:
    ratio = min(max(payload_kg / vehicle.max_payload_kg, 0.0), 1.0)
    return vehicle.empty_range_m - (vehicle.empty_range_m - vehicle.full_range_m) * ratio ** 1.5


def leg_energy(vehicle: Vehicle, leg: Leg, payload_kg: float) -> tuple[float, float, float]:
    horizontal = vehicle.usable_energy_kwh * leg.horizontal_distance_m / equivalent_range_m(vehicle, payload_kg)
    climb = (vehicle.empty_mass_kg + payload_kg) * G * leg.climb_m / (3_600_000.0 * vehicle.climb_efficiency)
    return horizontal + climb, horizontal, climb


def leg_flight_time_s(vehicle: Vehicle, leg: Leg) -> float:
    return (
        leg.climb_m / vehicle.climb_speed_mps
        + leg.horizontal_distance_m / vehicle.cruise_speed_mps
        + leg.descent_m / vehicle.descent_speed_mps
    )


def trip_metrics(vehicle: Vehicle, outbound: Leg, inbound: Leg, payload_kg: float, volume_m3: float, box_count: int, reserve_fraction: float) -> dict[str, Any]:
    e_out, e_out_h, e_out_up = leg_energy(vehicle, outbound, payload_kg)
    e_back, e_back_h, e_back_up = leg_energy(vehicle, inbound, 0.0)
    total_energy = e_out + e_back
    allowed = (1.0 - reserve_fraction) * vehicle.usable_energy_kwh
    time_s = (
        vehicle.fixed_prep_s
        + box_count * vehicle.per_box_load_s
        + leg_flight_time_s(vehicle, outbound)
        + vehicle.base_handoff_s
        + box_count * vehicle.per_box_handoff_s
        + leg_flight_time_s(vehicle, inbound)
    )
    mass_ok = payload_kg <= vehicle.max_payload_kg + 1e-9
    volume_ok = volume_m3 <= vehicle.max_volume_m3 + 1e-9
    energy_ok = total_energy <= allowed + 1e-9
    margin = allowed - total_energy
    if abs(margin) < 1e-8:
        margin = 0.0
    return {
        "feasible": bool(mass_ok and volume_ok and energy_ok),
        "mass_ok": bool(mass_ok),
        "volume_ok": bool(volume_ok),
        "energy_ok": bool(energy_ok),
        "energy_outbound_kwh": e_out,
        "energy_outbound_horizontal_kwh": e_out_h,
        "energy_outbound_climb_kwh": e_out_up,
        "energy_return_empty_kwh": e_back,
        "energy_return_horizontal_kwh": e_back_h,
        "energy_return_climb_kwh": e_back_up,
        "energy_total_kwh": total_energy,
        "allowed_energy_kwh": allowed,
        "energy_margin_kwh": margin,
        "soc_end": 1.0 - total_energy / vehicle.usable_energy_kwh,
        "operation_time_s": time_s,
    }


def max_safe_payload(vehicle: Vehicle, outbound: Leg, inbound: Leg, reserve_fraction: float) -> tuple[float | None, dict[str, Any] | None, str]:
    zero = trip_metrics(vehicle, outbound, inbound, 0.0, 0.0, 0, reserve_fraction)
    if zero["energy_total_kwh"] > zero["allowed_energy_kwh"]:
        return None, None, "zero_payload_energy_infeasible"
    full = trip_metrics(vehicle, outbound, inbound, vehicle.max_payload_kg, 0.0, 0, reserve_fraction)
    if full["energy_total_kwh"] <= full["allowed_energy_kwh"]:
        return vehicle.max_payload_kg, full, "rated_payload"
    low, high = 0.0, vehicle.max_payload_kg
    for _ in range(60):
        mid = (low + high) / 2.0
        metrics = trip_metrics(vehicle, outbound, inbound, mid, 0.0, 0, reserve_fraction)
        if metrics["energy_total_kwh"] <= metrics["allowed_energy_kwh"]:
            low = mid
        else:
            high = mid
    metrics = trip_metrics(vehicle, outbound, inbound, low, 0.0, 0, reserve_fraction)
    return low, metrics, "energy"


def add_score(a: tuple[float, float, float], b: tuple[float, float, float]) -> tuple[float, float, float]:
    return a[0] + b[0], a[1] + b[1], a[2] + b[2]


def ranking(score: tuple[float, float, float], objective: str) -> tuple[float, float, float]:
    flights, energy, duration = score
    if objective == "flights_energy_time":
        return flights, energy, duration
    if objective == "energy_flights_time":
        return energy, flights, duration
    if objective == "time_flights_energy":
        return duration, flights, energy
    raise ValueError(f"Unknown objective {objective}")


def solve_service(
    boxes: list[Box],
    vehicles: dict[str, Vehicle],
    outbound: Leg,
    inbound: Leg,
    reserve_fraction: float,
    objective: str,
) -> tuple[tuple[float, float, float], list[dict[str, Any]]]:
    n = len(boxes)
    full_mask = (1 << n) - 1
    masses = np.zeros(full_mask + 1)
    volumes = np.zeros(full_mask + 1)
    counts = np.zeros(full_mask + 1, dtype=int)
    feasible: list[list[tuple[str, dict[str, Any]]]] = [[] for _ in range(full_mask + 1)]
    for mask in range(1, full_mask + 1):
        bit = mask & -mask
        idx = bit.bit_length() - 1
        old = mask ^ bit
        masses[mask] = masses[old] + boxes[idx].mass_kg
        volumes[mask] = volumes[old] + boxes[idx].volume_m3
        counts[mask] = counts[old] + 1
        for vehicle_type in sorted(vehicles):
            vehicle = vehicles[vehicle_type]
            metrics = trip_metrics(
                vehicle,
                outbound,
                inbound,
                float(masses[mask]),
                float(volumes[mask]),
                int(counts[mask]),
                reserve_fraction,
            )
            if metrics["feasible"]:
                feasible[mask].append((vehicle_type, metrics))

    @lru_cache(maxsize=None)
    def dp(remaining: int) -> tuple[tuple[float, float, float], tuple[tuple[int, str], ...]] | None:
        if remaining == 0:
            return (0.0, 0.0, 0.0), ()
        anchor = remaining & -remaining
        subset = remaining
        best: tuple[tuple[float, float, float], tuple[tuple[int, str], ...]] | None = None
        while subset:
            if subset & anchor and feasible[subset]:
                rest = dp(remaining ^ subset)
                if rest is not None:
                    for vehicle_type, metrics in feasible[subset]:
                        score = add_score(
                            rest[0],
                            (1.0, float(metrics["energy_total_kwh"]), float(metrics["operation_time_s"])),
                        )
                        plan = ((subset, vehicle_type),) + rest[1]
                        candidate = (score, plan)
                        if best is None:
                            best = candidate
                        else:
                            cr, br = ranking(candidate[0], objective), ranking(best[0], objective)
                            if cr < br or (cr == br and candidate[1] < best[1]):
                                best = candidate
            subset = (subset - 1) & remaining
        return best

    solved = dp(full_mask)
    if solved is None:
        raise ValueError(f"No feasible partition for {boxes[0].service_node}")
    score, raw_plan = solved
    plan: list[dict[str, Any]] = []
    for mask, vehicle_type in raw_plan:
        vehicle = vehicles[vehicle_type]
        metrics = trip_metrics(
            vehicle,
            outbound,
            inbound,
            float(masses[mask]),
            float(volumes[mask]),
            int(counts[mask]),
            reserve_fraction,
        )
        plan.append(
            {
                "service_node": boxes[0].service_node,
                "vehicle_type": vehicle_type,
                "box_ids": [boxes[i].box_id for i in range(n) if mask & (1 << i)],
                "batch_total_mass_kg": float(masses[mask]),
                "batch_total_volume_m3": float(volumes[mask]),
                "box_count": int(counts[mask]),
                **metrics,
            }
        )
    plan.sort(key=lambda p: (p["vehicle_type"], p["box_ids"]))
    return score, plan


def solve_all_services(
    boxes_by_service: dict[str, list[Box]],
    vehicles: dict[str, Vehicle],
    legs: dict[tuple[str, str], Leg],
    reserve_fraction: float,
    objective: str,
) -> tuple[dict[str, tuple[float, float, float]], list[dict[str, Any]]]:
    service_scores: dict[str, tuple[float, float, float]] = {}
    all_batches: list[dict[str, Any]] = []
    for service in sorted(boxes_by_service):
        score, batches = solve_service(
            boxes_by_service[service],
            vehicles,
            legs[("O01", service)],
            legs[(service, "O01")],
            reserve_fraction,
            objective,
        )
        service_scores[service] = score
        all_batches.extend(batches)
    for i, batch in enumerate(all_batches, start=1):
        batch["batch_id"] = f"Q1-B{i:03d}"
    return service_scores, all_batches


def audit_batches(
    batches: list[dict[str, Any]],
    boxes_by_service: dict[str, list[Box]],
    vehicles: dict[str, Vehicle],
    legs: dict[tuple[str, str], Leg],
    reserve_fraction: float,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    box_map = {box.box_id: box for group in boxes_by_service.values() for box in group}
    assigned = [box_id for batch in batches for box_id in batch["box_ids"]]
    counts = Counter(assigned)
    audit_rows: list[dict[str, Any]] = []
    for batch in batches:
        service = batch["service_node"]
        vehicle = vehicles[batch["vehicle_type"]]
        listed = [box_map[box_id] for box_id in batch["box_ids"]]
        mass = sum(box.mass_kg for box in listed)
        volume = sum(box.volume_m3 for box in listed)
        metrics = trip_metrics(
            vehicle,
            legs[("O01", service)],
            legs[(service, "O01")],
            mass,
            volume,
            len(listed),
            reserve_fraction,
        )
        same_service = all(box.service_node == service for box in listed)
        values_match = (
            abs(mass - batch["batch_total_mass_kg"]) <= 1e-8
            and abs(volume - batch["batch_total_volume_m3"]) <= 1e-8
            and abs(metrics["energy_total_kwh"] - batch["energy_total_kwh"]) <= 1e-8
            and abs(metrics["operation_time_s"] - batch["operation_time_s"]) <= 1e-8
        )
        audit_rows.append(
            {
                "batch_id": batch["batch_id"],
                "service_node": service,
                "vehicle_type": vehicle.vehicle_type,
                "unique_within_batch": len(batch["box_ids"]) == len(set(batch["box_ids"])),
                "same_service_ok": same_service,
                "mass_ok": metrics["mass_ok"],
                "volume_ok": metrics["volume_ok"],
                "energy_ok": metrics["energy_ok"],
                "values_recomputed_ok": values_match,
                "energy_margin_kwh": metrics["energy_margin_kwh"],
                "all_constraints_ok": bool(
                    len(batch["box_ids"]) == len(set(batch["box_ids"]))
                    and same_service
                    and metrics["mass_ok"]
                    and metrics["volume_ok"]
                    and metrics["energy_ok"]
                    and values_match
                ),
            }
        )
    expected = set(box_map)
    global_audit = {
        "expected_box_count": len(expected),
        "assigned_box_count": len(assigned),
        "assigned_unique_box_count": len(set(assigned)),
        "missing_boxes": sorted(expected - set(assigned)),
        "unexpected_boxes": sorted(set(assigned) - expected),
        "duplicate_boxes": sorted(box_id for box_id, count in counts.items() if count != 1),
        "every_box_exactly_once": expected == set(assigned) and all(count == 1 for count in counts.values()),
        "all_batches_feasible": all(row["all_constraints_ok"] for row in audit_rows),
    }
    global_audit["overall_pass"] = global_audit["every_box_exactly_once"] and global_audit["all_batches_feasible"]
    return audit_rows, global_audit


def write_csv(path: Path, rows: Iterable[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    rows = list(rows)
    if not rows and fieldnames is None:
        raise ValueError(f"Cannot infer columns for empty output {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = fieldnames or list(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def workbook_from_csvs(output_path: Path, csv_paths: list[Path]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    alt_fill = PatternFill("solid", fgColor="D9EAF7")
    for csv_path in csv_paths:
        title = csv_path.stem[:31]
        ws = wb.create_sheet(title)
        with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(r_idx, c_idx)
                if r_idx == 1:
                    cell.value = value
                    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    try:
                        if value.lower() in {"true", "false"}:
                            cell.value = value.lower() == "true"
                        elif value != "" and all(ch not in value for ch in ["-", ":", "T"]):
                            cell.value = float(value)
                        else:
                            cell.value = value
                    except (ValueError, AttributeError):
                        cell.value = value
                    cell.font = Font(name="Arial", color="000000")
                    if r_idx % 2 == 0:
                        cell.fill = alt_fill
        if rows:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                letter = col[0].column_letter
                width = min(36, max(10, max(len(str(cell.value or "")) for cell in col) + 2))
                ws.column_dimensions[letter].width = width
    wb.save(output_path)


def plot_safe_payload(safe_rows: list[dict[str, Any]], output: Path) -> None:
    services = sorted({row["service_node"] for row in safe_rows})
    vehicles = ["A", "B", "C"]
    lookup = {(row["service_node"], row["vehicle_type"]): row["max_safe_payload_kg"] for row in safe_rows}
    matrix = np.array([[lookup[(s, v)] if lookup[(s, v)] is not None else np.nan for s in services] for v in vehicles])
    cmap = LinearSegmentedColormap.from_list("payload", ["#f7fbff", "#6baed6", "#08306b"])
    fig, ax = plt.subplots(figsize=(13, 3.8))
    im = ax.imshow(matrix, aspect="auto", cmap=cmap)
    ax.set_xticks(range(len(services)), services, rotation=45, ha="right")
    ax.set_yticks(range(len(vehicles)), vehicles)
    ax.set_xlabel("Service node")
    ax.set_ylabel("Vehicle type")
    ax.set_title("Maximum safe payload at the default 20% reserve")
    for i in range(len(vehicles)):
        for j in range(len(services)):
            text = "NA" if np.isnan(matrix[i, j]) else f"{matrix[i, j]:.1f}"
            ax.text(j, i, text, ha="center", va="center", fontsize=7, color="white" if matrix[i, j] > 45 else "black")
    fig.colorbar(im, ax=ax, label="kg", fraction=0.025, pad=0.02)
    fig.tight_layout()
    fig.savefig(output, dpi=220)
    plt.close(fig)


def plot_sensitivity(summary_rows: list[dict[str, Any]], output: Path) -> None:
    x = [100 * row["reserve_fraction"] for row in summary_rows]
    flights = [row["flights"] for row in summary_rows]
    energy = [row["energy_kwh"] for row in summary_rows]
    fig, ax1 = plt.subplots(figsize=(8.5, 4.8))
    ax1.plot(x, flights, marker="o", color="#1f77b4", label="Flights")
    ax1.set_xlabel("Return-energy reserve (%)")
    ax1.set_ylabel("Flights", color="#1f77b4")
    ax1.tick_params(axis="y", labelcolor="#1f77b4")
    ax1.grid(True, alpha=0.25)
    ax2 = ax1.twinx()
    ax2.plot(x, energy, marker="s", color="#d62728", label="Energy")
    ax2.set_ylabel("Total energy (kWh)", color="#d62728")
    ax2.tick_params(axis="y", labelcolor="#d62728")
    ax1.set_title("Q1 sensitivity to return-energy reserve")
    fig.tight_layout()
    fig.savefig(output, dpi=220)
    plt.close(fig)


def run(data_root: Path, output_dir: Path, config_path: Path) -> None:
    started = time.perf_counter()
    config = json.loads(config_path.read_text(encoding="utf-8"))
    default_reserve = float(config["default_reserve_fraction"])
    sensitivity_reserves = [float(x) for x in config["sensitivity_reserve_fractions"]]
    objectives = list(config["tradeoff_objectives"])

    nodes, vehicles, boxes_by_service, dem, input_files = load_inputs(data_root)
    if len(nodes) != 16 or sum(map(len, boxes_by_service.values())) != 80:
        raise ValueError("Input integrity check failed: expected 16 nodes and 80 boxes")

    legs: dict[tuple[str, str], Leg] = {}
    depot = nodes["O01"]
    for service in sorted(boxes_by_service):
        target = nodes[service]
        legs[("O01", service)] = build_leg(
            depot,
            target,
            depot.ground_elevation_m,
            target.ground_elevation_m + 30.0,
            dem,
        )
        legs[(service, "O01")] = build_leg(
            target,
            depot,
            target.ground_elevation_m + 30.0,
            depot.ground_elevation_m,
            dem,
        )

    node_rows = []
    for node_id in sorted(nodes, key=lambda value: (value != "O01", value)):
        node = nodes[node_id]
        x, y = local_xy_m(node, depot)
        node_rows.append(
            {
                "node_id": node.node_id,
                "node_type": node.node_type,
                "x": x,
                "y": y,
                "ground_elevation_m": node.ground_elevation_m,
                "longitude_deg": node.lon,
                "latitude_deg": node.lat,
            }
        )
    write_csv(output_dir / "node_registry.csv", node_rows)

    segment_rows = []
    for (from_node, to_node), leg in sorted(legs.items()):
        for vehicle_type, vehicle in sorted(vehicles.items()):
            segment_rows.append(
                {
                    "from_node": from_node,
                    "to_node": to_node,
                    "vehicle_type": vehicle_type,
                    "horizontal_distance_m": leg.horizontal_distance_m,
                    "cruise_altitude_m": leg.cruise_altitude_m,
                    "climb_m": leg.climb_m,
                    "descent_m": leg.descent_m,
                    "flight_time_s": leg_flight_time_s(vehicle, leg),
                    "energy_model_version": ENERGY_MODEL_VERSION,
                    "max_dem_elevation_m": leg.max_dem_elevation_m,
                    "intersected_dem_cells": leg.intersected_cells,
                    "geometry_model_version": GEOMETRY_MODEL_VERSION,
                }
            )
    write_csv(output_dir / "segment_library.csv", segment_rows)

    safe_rows = []
    for service in sorted(boxes_by_service):
        for vehicle_type, vehicle in sorted(vehicles.items()):
            payload, metrics, limiting = max_safe_payload(
                vehicle,
                legs[("O01", service)],
                legs[(service, "O01")],
                default_reserve,
            )
            safe_rows.append(
                {
                    "service_node": service,
                    "vehicle_type": vehicle_type,
                    "max_safe_payload_kg": payload,
                    "energy_outbound_kwh": None if metrics is None else metrics["energy_outbound_kwh"],
                    "energy_return_empty_kwh": None if metrics is None else metrics["energy_return_empty_kwh"],
                    "energy_total_kwh": None if metrics is None else metrics["energy_total_kwh"],
                    "energy_margin_kwh": None if metrics is None else metrics["energy_margin_kwh"],
                    "limiting_constraint": limiting,
                }
            )
    write_csv(output_dir / "safe_payload.csv", safe_rows)

    tradeoff_rows = []
    solution_by_objective: dict[str, list[dict[str, Any]]] = {}
    default_scores: dict[str, dict[str, tuple[float, float, float]]] = {}
    for objective in objectives:
        scores, batches = solve_all_services(
            boxes_by_service, vehicles, legs, default_reserve, objective
        )
        solution_by_objective[objective] = batches
        default_scores[objective] = scores
        tradeoff_rows.append(
            {
                "objective_priority": objective,
                "flights": len(batches),
                "energy_kwh": sum(batch["energy_total_kwh"] for batch in batches),
                "cumulative_operation_time_s": sum(batch["operation_time_s"] for batch in batches),
                "delivered_boxes": sum(batch["box_count"] for batch in batches),
            }
        )
    write_csv(output_dir / "tradeoff_summary.csv", tradeoff_rows)

    primary_objective = objectives[0]
    primary_batches = solution_by_objective[primary_objective]
    batching_rows = []
    batch_summary_rows = []
    box_lookup = {box.box_id: box for group in boxes_by_service.values() for box in group}
    for batch in primary_batches:
        batch_summary_rows.append(
            {
                "batch_id": batch["batch_id"],
                "service_node": batch["service_node"],
                "vehicle_type": batch["vehicle_type"],
                "box_count": batch["box_count"],
                "batch_total_mass_kg": batch["batch_total_mass_kg"],
                "batch_total_volume_m3": batch["batch_total_volume_m3"],
                "batch_energy_kwh": batch["energy_total_kwh"],
                "energy_margin_kwh": batch["energy_margin_kwh"],
                "soc_end": batch["soc_end"],
                "operation_time_s": batch["operation_time_s"],
            }
        )
        for box_id in batch["box_ids"]:
            box = box_lookup[box_id]
            batching_rows.append(
                {
                    "batch_id": batch["batch_id"],
                    "service_node": batch["service_node"],
                    "vehicle_type": batch["vehicle_type"],
                    "box_id": box.box_id,
                    "box_mass_kg": box.mass_kg,
                    "box_volume_m3": box.volume_m3,
                    "batch_total_mass_kg": batch["batch_total_mass_kg"],
                    "batch_total_volume_m3": batch["batch_total_volume_m3"],
                    "batch_energy_kwh": batch["energy_total_kwh"],
                    "batch_operation_time_s": batch["operation_time_s"],
                    "energy_margin_kwh": batch["energy_margin_kwh"],
                    "soc_end": batch["soc_end"],
                }
            )
    write_csv(output_dir / "batching_baseline.csv", batching_rows)
    write_csv(output_dir / "batch_summary.csv", batch_summary_rows)

    audit_rows, global_audit = audit_batches(
        primary_batches, boxes_by_service, vehicles, legs, default_reserve
    )
    write_csv(output_dir / "constraint_audit.csv", audit_rows)
    (output_dir / "global_audit.json").write_text(
        json.dumps(global_audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if not global_audit["overall_pass"]:
        raise AssertionError("Primary Q1 solution failed its audit")

    sensitivity_payload_rows = []
    sensitivity_summary_rows = []
    for reserve in sensitivity_reserves:
        for service in sorted(boxes_by_service):
            for vehicle_type, vehicle in sorted(vehicles.items()):
                payload, metrics, limiting = max_safe_payload(
                    vehicle,
                    legs[("O01", service)],
                    legs[(service, "O01")],
                    reserve,
                )
                sensitivity_payload_rows.append(
                    {
                        "reserve_fraction": reserve,
                        "service_node": service,
                        "vehicle_type": vehicle_type,
                        "max_safe_payload_kg": payload,
                        "limiting_constraint": limiting,
                        "energy_margin_kwh": None if metrics is None else metrics["energy_margin_kwh"],
                    }
                )
        try:
            _, batches = solve_all_services(
                boxes_by_service, vehicles, legs, reserve, primary_objective
            )
            sensitivity_summary_rows.append(
                {
                    "reserve_fraction": reserve,
                    "feasible": True,
                    "flights": len(batches),
                    "energy_kwh": sum(batch["energy_total_kwh"] for batch in batches),
                    "cumulative_operation_time_s": sum(batch["operation_time_s"] for batch in batches),
                }
            )
        except ValueError:
            sensitivity_summary_rows.append(
                {
                    "reserve_fraction": reserve,
                    "feasible": False,
                    "flights": None,
                    "energy_kwh": None,
                    "cumulative_operation_time_s": None,
                }
            )
    write_csv(output_dir / "safe_payload_sensitivity.csv", sensitivity_payload_rows)
    write_csv(output_dir / "sensitivity_summary.csv", sensitivity_summary_rows)

    service_summary_rows = []
    primary_scores = default_scores[primary_objective]
    for service, score in sorted(primary_scores.items()):
        service_batches = [b for b in primary_batches if b["service_node"] == service]
        service_summary_rows.append(
            {
                "service_node": service,
                "box_count": len(boxes_by_service[service]),
                "total_mass_kg": sum(box.mass_kg for box in boxes_by_service[service]),
                "total_volume_m3": sum(box.volume_m3 for box in boxes_by_service[service]),
                "flights": int(score[0]),
                "energy_kwh": score[1],
                "cumulative_operation_time_s": score[2],
                "vehicle_mix": ";".join(f"{k}:{v}" for k, v in sorted(Counter(b["vehicle_type"] for b in service_batches).items())),
            }
        )
    write_csv(output_dir / "service_summary.csv", service_summary_rows)

    plot_safe_payload(safe_rows, output_dir / "safe_payload_heatmap.png")
    feasible_sensitivity = [row for row in sensitivity_summary_rows if row["feasible"]]
    plot_sensitivity(feasible_sensitivity, output_dir / "reserve_sensitivity.png")

    csv_paths = [
        output_dir / "safe_payload.csv",
        output_dir / "batching_baseline.csv",
        output_dir / "batch_summary.csv",
        output_dir / "service_summary.csv",
        output_dir / "tradeoff_summary.csv",
        output_dir / "sensitivity_summary.csv",
        output_dir / "constraint_audit.csv",
        output_dir / "segment_library.csv",
        output_dir / "node_registry.csv",
    ]
    workbook_from_csvs(output_dir / "Q1_results.xlsx", csv_paths)

    elapsed = time.perf_counter() - started
    primary_tradeoff = next(row for row in tradeoff_rows if row["objective_priority"] == primary_objective)
    summary = {
        "status": "integrable" if global_audit["overall_pass"] else "audit_failed",
        "primary_objective": primary_objective,
        "default_reserve_fraction": default_reserve,
        "delivered_boxes": primary_tradeoff["delivered_boxes"],
        "flights": primary_tradeoff["flights"],
        "energy_kwh": primary_tradeoff["energy_kwh"],
        "cumulative_operation_time_s": primary_tradeoff["cumulative_operation_time_s"],
        "audit_pass": global_audit["overall_pass"],
        "energy_model_version": ENERGY_MODEL_VERSION,
        "geometry_model_version": GEOMETRY_MODEL_VERSION,
        "elapsed_seconds": elapsed,
    }
    (output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    manifest = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "python": sys.version,
        "platform": platform.platform(),
        "command": "python code/q1_solver.py --data-root <problem_data_root> --output-dir results",
        "deterministic": True,
        "random_seed": None,
        "config": config,
        "input_sha256": {path.name: sha256_file(path) for path in input_files},
        "code_sha256": sha256_file(Path(__file__)),
        "elapsed_seconds": elapsed,
    }
    (output_dir / "runtime_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Exact Q1 single-service round-trip batching solver")
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument(
        "--config",
        type=Path,
        default=Path(__file__).with_name("config.json"),
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    run(args.data_root.resolve(), args.output_dir.resolve(), args.config.resolve())
