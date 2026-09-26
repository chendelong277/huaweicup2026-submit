from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
OUT = ROOT / "Q1_实验报告.docx"


def read_csv(name: str) -> list[dict[str, str]]:
    with (RESULTS / name).open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_picture_with_alt(paragraph, image_path, width, title, description):
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=width)
    inline = run._r.xpath(".//wp:inline")
    if inline:
        doc_pr = inline[0].find(qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("title", title)
            doc_pr.set("descr", description)


def set_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(p, before=0, after=6, line=1.25, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_text(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    style_paragraph(p, before=10 if level == 1 else 6, after=4, line=1.1)
    r = p.add_run(text)
    set_font(r, size=14 if level == 1 else 11.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        r = p.add_run(str(value))
        set_font(r, size=font_size, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if ridx % 2 == 1:
                set_cell_shading(cell, "F3F7FA")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.0)
            r = p.add_run(str(value))
            set_font(r, size=font_size)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def fmt(v, digits=3):
    if v in (None, "", "nan"):
        return "—"
    try:
        return f"{float(v):.{digits}f}"
    except Exception:
        return str(v)


def main():
    summary = json.loads((RESULTS / "summary.json").read_text(encoding="utf-8"))
    audit = json.loads((RESULTS / "global_audit.json").read_text(encoding="utf-8"))
    safe = read_csv("safe_payload.csv")
    batches = read_csv("batch_summary.csv")
    service = read_csv("service_summary.csv")
    tradeoff = read_csv("tradeoff_summary.csv")
    sensitivity = read_csv("sensitivity_summary.csv")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].font.name = "Microsoft YaHei"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    title_style = styles["Title"]
    title_style.font.name = "Microsoft YaHei"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title_style._element.pPr
    if title_ppr is not None:
        title_bdr = title_ppr.find(qn("w:pBdr"))
        if title_bdr is not None:
            title_ppr.remove(title_bdr)

    # Header/footer
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("D题问题一实验报告")
    set_font(hr, size=8.5, color="666666")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Q1 单点往返运输能力与货箱组批")
    set_font(fr, size=8.5, color="666666")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, before=18, after=8, line=1.1)
    r = title.add_run("问题一 单点往返运输能力与货箱组批实验报告")
    set_font(r, size=20, bold=True)
    remove_paragraph_borders(title)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(sub, after=16, line=1.1)
    r = sub.add_run("山区洪涝灾害下无人机运输与通信协同优化")
    set_font(r, size=11.5, color="555555")

    add_text(doc, "摘要：本实验针对问题一的单点直接往返运输任务，基于题目提供的节点坐标、30 m DEM、三类运输无人机参数和80个不可拆分货箱，计算15个服务区对应的最大安全载荷，并在质量、体积和返航安全能量余量约束下求解服务区内货箱组批。实验采用DEM直线穿越像元的保守几何判定、载荷相关等效航程能耗模型，以及按架次数、总能耗、累计作业时间的词典序精确动态规划。默认返航安全余量20%时，80个货箱被唯一配送，得到18架次，总运输能耗59.263908 kWh，累计作业时间32804.873 s；全部批次通过独立审计。敏感性实验显示，返航余量提高到25%、30%、35%时架次分别增至19、20、25，40%时整体任务不可行。")
    add_text(doc, "关键词：无人机运输；单点往返；货箱组批；DEM；返航安全余量；动态规划")

    add_heading(doc, "1 实验目标与问题定义", 1)
    add_text(doc, "问题一研究每个服务区独立由调度中心 O01 直接往返的运输能力。每个架次采用 O01→Si→O01，仅服务一个服务区；同一服务区允许多个架次，但货箱不可拆分、每个货箱只能安排一次，且不得跨服务区组批。实验需要同时给出三类机型在各服务区的最大安全载荷、可执行的货箱组批方案，以及架次数、能耗和累计作业时间之间的权衡。")
    add_text(doc, "本实验不考虑实体无人机数量、共享电池周转和跨服务区多点路径，这些约束属于问题二及后续问题。Q1 的组批结果作为后续问题的单点运输能力基线和候选任务包来源，而不是问题二、问题三的冻结解。")

    add_heading(doc, "2 数据与预处理", 1)
    add_text(doc, "实验读取题目目录中的四类输入：调度中心与服务区.xlsx、运输无人机数据.xlsx、物资需求与配送时限.xlsx，以及镇龙乡及周边30米DEM.mat。数据检查得到1个调度中心、15个服务区、3种运输机型和80个唯一货箱。质量统一使用 kg，体积使用 m³，距离使用 m，时间使用 s，能量使用 kWh，SOC 使用 0 至 1 的小数。")
    add_text(doc, "空间预处理以 O01 为参考建立局部平面米制坐标。对 O01 与每个服务区的水平直线，采用栅格 supercover 遍历，将直线相交或触及的DEM像元全部纳入，取其中最高地面高程并加50 m作为巡航海拔。O01 作业高度取地面海拔，服务区作业高度取地面海拔加30 m。所有服务航段均未穿过 DEM NoData 像元。")
    add_table(doc, ["数据对象", "数量", "关键字段", "用途"], [
        ["节点", "16", "编号、经纬度、海拔", "航段几何和作业高度"],
        ["运输机型", "3", "载重、体积、速度、航程、电量", "安全载荷和能耗"],
        ["货箱", "80", "服务区、质量、体积、类别", "不可拆分组批"],
        ["DEM", "30 m栅格", "高程、NoData、经纬度网格", "巡航海拔判定"],
    ], widths=[2.5, 2.0, 6.0, 5.0], font_size=8.8)

    add_heading(doc, "3 模型与求解方法", 1)
    add_heading(doc, "3.1 航段时间与能耗", 2)
    add_text(doc, "机型 g 携带载荷 q 时的等效航程按题面公式计算：L_g(q)=L_g0−(L_g0−L_gF)(q/Q_g)^(3/2)。水平巡航能耗采用 E_hor=E_use·d/L_g(q)，爬升附加能耗采用 E_up=(m_g+q)g h_up/(3.6×10^6 η_g)，下降附加能耗按题面取0。单点往返架次的去程载荷为批次总质量，返程载荷为0。架次时间由固定准备、逐箱装载、去程飞行、基础交接、逐箱交接和返程飞行构成。")
    add_heading(doc, "3.2 最大安全载荷", 2)
    add_text(doc, "对每个机型—服务区组合，最大安全载荷 q* 是满足 q* 不超过额定载荷且往返总能耗不超过 (1−ρ)E_use 的最大质量。由于能耗随载荷单调增加，程序使用二分搜索求解；货舱体积不并入连续质量上限，而在具体货箱组批时作为独立硬约束。")
    add_heading(doc, "3.3 货箱组批", 2)
    add_text(doc, "对每个服务区枚举全部非空货箱子集，并分别检查三种机型的质量、体积和往返能量约束，将通过检查的“货箱子集—机型”作为可行批次模式。随后以位掩码动态规划进行集合划分，固定剩余货箱中的最低位货箱以消除批次排列重复。主目标为词典序最小化（架次数，总能耗，累计作业时间）。由于单个服务区最多15箱，该方法能够完整枚举可行模式并得到主目标下的精确解。")
    add_heading(doc, "3.4 实验设置", 2)
    add_table(doc, ["设置项", "取值"], [
        ["默认返航安全余量", "20%"],
        ["主目标优先序", "架次数 → 总能耗 → 累计作业时间"],
        ["对比目标", "能耗 → 架次数 → 时间；时间 → 架次数 → 能耗"],
        ["敏感性场景", "10%、15%、20%、25%、30%、35%、40%"],
        ["随机性", "无随机数，确定性求解"],
    ], widths=[6.0, 9.5], font_size=8.8)

    add_heading(doc, "4 最大安全载荷结果", 1)
    add_text(doc, "默认返航安全余量20%时，A型在全部服务区均受额定载荷25 kg限制；B型除S008外均为额定载荷30 kg，S008受能量限制为28.630 kg；C型在S001、S005-S007、S009-S011及S013-S015可达到额定载荷80 kg，在S002、S003、S004、S008、S012受能量限制。")
    safe_rows = []
    for service_node in sorted({row["service_node"] for row in safe}):
        lookup = {row["vehicle_type"]: row for row in safe if row["service_node"] == service_node}
        safe_rows.append([service_node, fmt(lookup["A"]["max_safe_payload_kg"], 3), fmt(lookup["B"]["max_safe_payload_kg"], 3), fmt(lookup["C"]["max_safe_payload_kg"], 3), "; ".join(f"{k}:{lookup[k]['limiting_constraint']}" for k in ("A", "B", "C"))])
    add_table(doc, ["服务区", "A型 kg", "B型 kg", "C型 kg", "限制来源"], safe_rows, widths=[2.3, 2.5, 2.5, 2.5, 6.0], font_size=8.1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=2, after=2)
    add_picture_with_alt(p, RESULTS / "safe_payload_heatmap.png", Cm(16.2), "最大安全载荷热力图", "默认20%返航安全余量下，A、B、C三类运输机在15个服务区的最大安全载荷，单位为kg。")
    cap = doc.add_paragraph("图1  默认20%返航安全余量下的最大安全载荷（kg）")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(cap, after=8)
    set_font(cap.runs[0], size=8.5, color="555555")

    add_heading(doc, "5 货箱组批与目标权衡", 1)
    add_text(doc, "主方案共18个批次。S001、S002、S003分别需要2架次，其余服务区各1架次。S001采用2架C型；S002和S003各采用1架B型与1架C型；S004-S008采用C型；S009-S015采用B型。所有80个货箱均被安排一次。")
    add_table(doc, ["方案目标优先序", "架次", "总能耗 kWh", "累计作业时间 s", "说明"], [[
        row["objective_priority"], row["flights"], fmt(row["energy_kwh"], 6), fmt(row["cumulative_operation_time_s"], 3),
        "主方案" if row["objective_priority"] == "flights_energy_time" else ("能耗优先" if row["objective_priority"] == "energy_flights_time" else "时间优先")
    ] for row in tradeoff], widths=[4.3, 2.0, 3.2, 4.0, 3.0], font_size=8.4)
    add_text(doc, "主方案相对于能耗优先方案多消耗约0.098056 kWh，但减少1架次并缩短约1668.083 s的累计作业时间；时间优先方案与主方案一致。因此，在未给出人为权重的条件下，将减少起降和架次作为第一优先级具有可解释性。")
    doc.add_page_break()
    add_table(doc, ["批次", "服务区", "机型", "箱数", "质量 kg", "体积 m³", "能耗 kWh", "作业时间 s"], [[
        row["batch_id"], row["service_node"], row["vehicle_type"], row["box_count"], fmt(row["batch_total_mass_kg"], 1), fmt(row["batch_total_volume_m3"], 3), fmt(row["batch_energy_kwh"], 3), fmt(row["operation_time_s"], 1)
    ] for row in batches], widths=[2.2, 2.2, 1.5, 1.4, 2.1, 2.2, 2.3, 2.7], font_size=7.4)

    add_heading(doc, "6 返航安全余量敏感性", 1)
    add_text(doc, "提高返航安全余量会压缩可用于运输的能量上限，影响首先出现在距离较远或地形抬升较大的服务区。实验结果呈阶梯变化：10%至20%时保持18架次；25%时增加到19架次；30%时为20架次；35%时为25架次；40%时无法完成全部货箱。")
    add_table(doc, ["返航余量", "可行性", "架次", "总能耗 kWh", "累计作业时间 s"], [[
        f"{float(row['reserve_fraction'])*100:.0f}%", "是" if row["feasible"] == "True" else "否", row["flights"] or "—", fmt(row["energy_kwh"], 6), fmt(row["cumulative_operation_time_s"], 3)
    ] for row in sensitivity], widths=[3.2, 3.0, 2.0, 4.0, 5.0], font_size=8.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, before=2, after=2)
    add_picture_with_alt(p, RESULTS / "reserve_sensitivity.png", Cm(16.0), "返航余量敏感性图", "返航安全余量从10%增加到35%时架次和总运输能耗的变化；40%场景不可行。")
    cap = doc.add_paragraph("图2  返航安全余量对架次和总能耗的影响")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(cap, after=8)
    set_font(cap.runs[0], size=8.5, color="555555")
    add_text(doc, "40%余量场景的不可行原因不是组批算法无法找到划分，而是S008的三种机型在空载往返时也无法满足返航能量上限。因此该场景应解释为系统能力边界，而不是通过继续拆分货箱来修复。")

    add_heading(doc, "7 约束审计与可复现性", 1)
    add_text(doc, f"独立审计结果：期望货箱数{audit['expected_box_count']}，已分配货箱数{audit['assigned_box_count']}，唯一货箱数{audit['assigned_unique_box_count']}，漏箱{len(audit['missing_boxes'])}，重复箱{len(audit['duplicate_boxes'])}；每个货箱恰好一次={str(audit['every_box_exactly_once']).lower()}，所有批次可行={str(audit['all_batches_feasible']).lower()}，总体审计通过={str(audit['overall_pass']).lower()}。")
    add_table(doc, ["审计项目", "结果"], [
        ["货箱唯一配送", "通过：80/80，漏箱0，重复0"],
        ["批次服务区约束", "通过：每个批次仅含一个服务区"],
        ["质量约束", "通过：18/18批次"],
        ["体积约束", "通过：18/18批次"],
        ["返航能量约束", "通过：18/18批次"],
        ["结果文件校验", "通过：CSV、Excel和审计JSON已生成"],
    ], widths=[6.0, 9.5], font_size=8.8)
    add_text(doc, "复现实验入口为 code/q1_solver.py，参数位于 code/config.json，输入文件和代码的SHA-256、运行环境及运行时间记录在 results/runtime_manifest.json。结果文件采用UTF-8 CSV；Q1_results.xlsx为便于阅读的汇总工作簿，未包含公式和外部链接。")

    add_heading(doc, "8 结论与适用范围", 1)
    add_text(doc, "在默认20%返航安全余量下，单点直接往返的Q1主方案能够以18架次完成80个货箱交付，且所有硬约束均满足。服务区距离与地形抬升使C型机在部分服务区由额定载荷转为能量受限，S008是高返航余量下最敏感的服务区。主方案适合作为Q2的单点可行基线和候选任务包来源。")
    add_text(doc, "本报告的能耗结果依赖当前对题面未展开能耗公式的解释：水平能耗按等效航程折算，爬升能耗按重力势能除以爬升效率计算。若官方补充材料给出新的显式公式，应在统一接口层替换能耗函数，并重新生成Q1至Q3的相关结果。")

    add_heading(doc, "附录 A 交付文件", 1)
    add_table(doc, ["文件", "作用"], [
        ["results/safe_payload.csv", "三机型—服务区最大安全载荷"],
        ["results/batching_baseline.csv", "逐箱组批接口"],
        ["results/segment_library.csv", "公共航段库"],
        ["results/node_registry.csv", "节点注册表"],
        ["results/constraint_audit.csv", "批次级约束审计"],
        ["results/global_audit.json", "全局唯一配送审计"],
        ["results/Q1_results.xlsx", "结果汇总工作簿"],
        ["code/q1_solver.py", "可复现求解程序"],
    ], widths=[7.0, 8.5], font_size=8.6)

    # Core properties avoid exposing personal metadata.
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "问题一 单点往返运输能力与货箱组批实验报告"
    doc.core_properties.subject = "山区洪涝灾害下无人机运输与通信协同优化"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
