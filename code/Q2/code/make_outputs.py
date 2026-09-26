from __future__ import annotations
import argparse, csv
from pathlib import Path
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def read(p):
    with p.open(encoding="utf-8-sig") as f:return list(csv.DictReader(f))

def main(results:Path, q1_results:Path):
    trips=read(results/"trip_summary.csv"); resources=read(results/"transport_resource_timeline.csv"); comparison=read(results/"plan_comparison.csv")
    nodes={r["node_id"]:(float(r["x"]),float(r["y"])) for r in read(q1_results/"node_registry.csv")}
    fig,ax=plt.subplots(figsize=(9,8))
    colors={"A":"#4c78a8","B":"#f58518","C":"#54a24b"}
    for t in trips:
        route=t["route"].split("->"); xy=[nodes[x] for x in route]
        ax.plot([p[0] for p in xy],[p[1] for p in xy],color=colors[t["vehicle_type"]],alpha=.35,lw=1.4)
    for n,(x,y) in nodes.items():
        ax.scatter(x,y,s=65 if n=="O01" else 28,c="#b22222" if n=="O01" else "#17384d",zorder=3); ax.text(x+80,y+80,n,fontsize=8)
    ax.set(title="Q2 transport routes",xlabel="Local x (m)",ylabel="Local y (m)"); ax.grid(alpha=.2); ax.set_aspect("equal",adjustable="datalim")
    ax.legend(handles=[Patch(color=colors[k],label=f"Type {k}") for k in ("B","C")],loc="best")
    fig.tight_layout(); fig.savefig(results/"route_map.png",dpi=220); plt.close(fig)

    usage=[r for r in resources if r["activity"] in {"flight_and_service","flight","charging"}]
    ids=sorted({r["resource_id"] for r in usage}); pos={x:i for i,x in enumerate(ids)}
    fig,ax=plt.subplots(figsize=(12,max(5,len(ids)*.38)))
    for r in usage:
        st=float(r["start_time_s"]); en=float(r["end_time_s"]); y=pos[r["resource_id"]]
        color="#4c78a8" if r["resource_type"]=="transport_uav" else ("#72b7b2" if r["activity"]=="flight" else "#f2a541")
        ax.barh(y,en-st,left=st,height=.55,color=color,alpha=.82)
    ax.set_yticks(range(len(ids)),ids); ax.set(xlabel="Time (s)",title="Q2 UAV and battery usage timeline"); ax.grid(axis="x",alpha=.2)
    ax.legend(handles=[Patch(color="#4c78a8",label="UAV occupied"),Patch(color="#72b7b2",label="Battery occupied"),Patch(color="#f2a541",label="Battery charging")],loc="lower right")
    fig.tight_layout(); fig.savefig(results/"resource_gantt.png",dpi=220); plt.close(fig)

    fig,axs=plt.subplots(1,3,figsize=(11,3.6)); names=[r["plan"].replace("_schedule","") for r in comparison]
    for ax,key,title in zip(axs,["hard_deadline_violations","makespan_s","energy_kwh"],["Hard deadline violations","Makespan (s)","Energy (kWh)"]):
        vals=[float(r[key]) for r in comparison]; ax.bar(names,vals,color=["#9ecae9","#31a354"]); ax.set_title(title); ax.tick_params(axis="x",rotation=18)
    fig.tight_layout(); fig.savefig(results/"plan_comparison.png",dpi=220); plt.close(fig)

    csvs=sorted(results.glob("*.csv")); wb=Workbook(); wb.remove(wb.active)
    for p in csvs:
        ws=wb.create_sheet(p.stem[:31]); rows=read(p)
        if not rows: continue
        headers=list(rows[0]); ws.append(headers)
        for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="D9EAF7")
        for r in rows: ws.append([r[h] for h in headers])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(45,max(10,max(len(str(x.value or "")) for x in col)+2))
    wb.save(results/"Q2_results.xlsx")

if __name__=="__main__":
    ap=argparse.ArgumentParser();ap.add_argument("--results-dir",type=Path,required=True);ap.add_argument("--q1-results-dir",type=Path,required=True);a=ap.parse_args();main(a.results_dir,a.q1_results_dir)
